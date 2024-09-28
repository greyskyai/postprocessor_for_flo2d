import os
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from modules.utilities import time_function
from modules.swmm_rating_tables_extraction import extract_swmm_rating_tables

@time_function
def plot_rating_tables_to_pdf(rating_tables, pdf_filename):
    """
    Plot rating tables (discharge vs stage) to a PDF file with 4 plots per page.

    Args:
    rating_tables (list): A list of dictionaries containing table names and data.
    pdf_filename (str): The output file path for the PDF.
    """
    with PdfPages(pdf_filename) as pdf:
        # Prepare the plot layout: 2x2 grid on each page
        num_plots_per_page = 4
        fig, axes = plt.subplots(2, 2, figsize=(8.5, 11))
        fig.subplots_adjust(hspace=0.3, wspace=0.3)
        
        # Track plot count
        plot_count = 0
        
        for idx, table in enumerate(rating_tables):
            ax = axes[plot_count // 2, plot_count % 2]
            ax.plot(table["Data"]["Flow"], table["Data"]["Stage"], color='blue', label='Stage vs Flow', marker='o')
            ax.set_title(f"Table: {table['Table']}")
            ax.set_xlabel('Discharge (cfs)')
            ax.set_ylabel('Stage (ft)')
            ax.grid(True)
            
            # Move to next plot position
            plot_count += 1
            
            # If we've filled the 2x2 grid, save the page and start a new one
            if plot_count == num_plots_per_page:
                pdf.savefig(fig)
                plt.close(fig)
                fig, axes = plt.subplots(2, 2, figsize=(8.5, 11))
                fig.subplots_adjust(hspace=0.3, wspace=0.3)
                plot_count = 0
        
        # Save any remaining plots on the final page
        if plot_count > 0:
            for remaining in range(plot_count, num_plots_per_page):
                axes[remaining // 2, remaining % 2].axis('off')  # Turn off unused subplots
            pdf.savefig(fig)
            plt.close(fig)

@time_function
def create_rating_tables_spreadsheet(rating_tables, excel_filename):
    """
    Create an Excel spreadsheet with rating table data and plots for each structure.
    The plots will include a blue line with blue data points.

    Args:
    rating_tables (list): A list of dictionaries containing table names and data.
    excel_filename (str): The output file path for the Excel spreadsheet.
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove the default sheet

    for table in rating_tables:
        table_name = table["Table"]
        data = table["Data"]

        # Create a new worksheet for each table
        ws = wb.create_sheet(title=table_name)

        # Write data to the worksheet
        ws.append(["Stage (ft)", "Flow (cfs)"])
        for _, row in data.iterrows():
            ws.append([row["Stage"], row["Flow"]])

        # Create a scatter plot
        chart = ScatterChart()
        chart.title = f"Rating Curve - {table_name}"
        chart.x_axis.title = "Flow (cfs)"
        chart.y_axis.title = "Stage (ft)"

        # Define data ranges for the chart
        x_values = Reference(ws, min_col=2, min_row=2, max_row=len(data) + 1)
        y_values = Reference(ws, min_col=1, min_row=2, max_row=len(data) + 1)

        # Create a single series for both line and points
        series = Series(y_values, x_values, title="Stage vs Flow")
        series.marker = Marker(symbol='circle', size=7)
        series.marker.graphicalProperties.solidFill = "4472C4"  # Blue color
        series.marker.graphicalProperties.line.solidFill = "4472C4"  # Blue outline
        series.graphicalProperties.line.solidFill = "4472C4"  # Blue line
        series.graphicalProperties.line.width = 20000  # Adjust line width (in EMUs)
        series.smooth = True  # This creates a smoothed line
        chart.series.append(series)

        # Add the chart to the worksheet
        ws.add_chart(chart, "D2")

    # Save the workbook
    wb.save(excel_filename)

@time_function
def swmm_rating_tables_and_plots(file_path, rating_tables):
    """
    Process SWMM rating tables: create plots and generate spreadsheets.

    Args:
    file_path (str): Path to the directory containing SWMMFLORT.DAT.
    rating_tables (list): List of dictionaries containing rating table data.

    Returns:
    tuple: Paths to the generated PDF and Excel files, or (None, None) if no data found.
    """
    output_folder = os.path.join(file_path, 'flo2d_plots')
    os.makedirs(output_folder, exist_ok=True)

    if rating_tables:
        pdf_filename = os.path.join(output_folder, 'swmm_rating_tables.pdf')
        excel_filename = os.path.join(output_folder, 'swmm_rating_tables.xlsx')

        plot_rating_tables_to_pdf(rating_tables, pdf_filename)
        create_rating_tables_spreadsheet(rating_tables, excel_filename)

        return pdf_filename, excel_filename
    else:
        return None, None

# If you want to test the functions when the script is run directly
if __name__ == "__main__":
    from swmm_rating_tables_extraction import extract_general_rating_tables
    
    folder_path = r"R:\_anichols\Projects\_flo2d_postprocessor_tests\Detroit_Basin_Prop100y24h"
    test_file_path = os.path.join(folder_path, "SWMMFLORT.DAT")  # Replace with an actual test file path
    rating_tables = extract_swmm_rating_tables(test_file_path)
    
    pdf_file, excel_file = swmm_rating_tables_and_plots(test_file_path, rating_tables)
    
    if pdf_file and excel_file:
        print(f"SWMM rating tables plotted and saved to: {pdf_file}")
        print(f"SWMM rating table spreadsheet created: {excel_file}")
    else:
        print("No SWMM rating tables found in SWMMFLORT.DAT")